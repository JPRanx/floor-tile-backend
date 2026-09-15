"""Deterministic adapters for Ashley's current production and dispatch files.

The adapters normalize source files into local preview/rehearsal carriers only.
Tentative dispatch rows are never promoted to in-transit authority.
"""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import PurePosixPath
from time import monotonic as _monotonic
from typing import Sequence
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

Q2 = Decimal("0.01")
_MAX_XLSX_ENTRIES = 512
_MAX_XLSX_TOTAL_BYTES = 64 * 1024 * 1024
_MAX_XLSX_ENTRY_BYTES = 32 * 1024 * 1024
_MAX_WORKSHEET_ROWS = 100_000
_MAX_WORKSHEET_COLUMNS = 128
_MAX_WORKSHEET_CELLS = 2_000_000
_MAX_PDF_PAGES = 100
_MAX_PDF_TEXT_BYTES = 2 * 1024 * 1024
_MAX_PDF_SECONDS = 15
_MAX_XLS_BYTES = 10 * 1024 * 1024
_MAX_XLS_SHEETS = 16


def _q2(value) -> Decimal:
    return Decimal(str(value)).quantize(Q2, rounding=ROUND_HALF_UP)


def _norm(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().upper()
    text = text.replace("(EA)", " ").replace("(T)", " ")
    text = re.sub(r"\b5[12][X*]5[123](?:-1)?\b", " ", text)
    text = text.replace("BTE", " ")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return " ".join(text.split())


def _catalog_map(catalog: Sequence[dict]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    collisions: set[str] = set()
    for product in catalog:
        canonical = str(product["sku"])
        values = [canonical, product.get("name"), *(product.get("aliases") or [])]
        for value in values:
            if not value:
                continue
            key = _norm(value)
            prior = mapping.get(key)
            if prior is not None and prior != canonical:
                collisions.add(key)
            else:
                mapping[key] = canonical
    for key in collisions:
        mapping.pop(key, None)
    return mapping


def _open_xlsx(file_bytes: bytes, *, data_only: bool = True):
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_XLSX_ENTRIES:
                raise ValueError
            total = 0
            for info in infos:
                normalized = info.filename.replace("\\", "/")
                path = PurePosixPath(normalized)
                if (info.flag_bits & 1 or path.is_absolute() or ".." in path.parts
                        or re.match(r"^[A-Za-z]:", normalized)):
                    raise ValueError
                if info.file_size > _MAX_XLSX_ENTRY_BYTES:
                    raise ValueError
                total += info.file_size
                if total > _MAX_XLSX_TOTAL_BYTES:
                    raise ValueError
        return load_workbook(BytesIO(file_bytes), read_only=True, data_only=data_only)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as exc:
        raise ValueError("invalid XLSX workbook") from exc


def _bounded_rows(sheet, **kwargs):
    max_row = sheet.max_row or 0
    max_column = sheet.max_column or 0
    if (max_row > _MAX_WORKSHEET_ROWS or max_column > _MAX_WORKSHEET_COLUMNS
            or max_row * max_column > _MAX_WORKSHEET_CELLS):
        raise ValueError("invalid XLSX workbook")
    visited = 0
    for row in sheet.iter_rows(**kwargs):
        visited += len(row)
        if visited > _MAX_WORKSHEET_CELLS:
            raise ValueError("invalid XLSX workbook")
        yield row


def _decimal(value, field: str) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"missing {field}")
    try:
        decimal = _q2(str(value).strip().replace(",", ""))
        if not decimal.is_finite():
            raise ValueError
        return decimal
    except (ValueError, ArithmeticError) as exc:
        raise ValueError(f"invalid {field}") from exc


def _identity(raw, mapping: dict[str, str]) -> str | None:
    if raw is None or not str(raw).strip():
        return None
    return mapping.get(_norm(raw)) or str(raw).strip()


def parse_warehouse_xlsx(file_bytes: bytes, *, catalog: Sequence[dict]) -> dict:
    """Normalize either accepted Tarragona closing-inventory layout."""
    workbook = _open_xlsx(file_bytes)
    try:
        sheet = workbook["Inventario"] if "Inventario" in workbook.sheetnames else workbook.active
        header_one = list(next(_bounded_rows(sheet, min_row=1, max_row=1, values_only=True)))
        header_two = list(next(_bounded_rows(sheet, min_row=2, max_row=2, values_only=True)))
        compact = (
            len(header_two) >= 5
            and [str(value or "").strip().upper() for value in header_two[:5]]
            == ["MARCA", "CATEGORÍA", "REFERENCIA", "FORMATO", "M²"]
        )
        grouped_headers = {
            0: "PRODUCTO", 4: "INVENTARIO INICIAL", 7: "INGRESOS",
            10: "SALIDAS", 13: "SALDO FINAL",
        }
        quantity_headers = (5, 8, 11, 14)
        detailed = not (
            len(header_one) <= 13 or len(header_two) <= 14
            or any(str(header_one[index] or "").strip().upper() != expected
                   for index, expected in grouped_headers.items())
            or str(header_two[2]).strip() != "Referencia"
            or any(str(header_two[index] or "").strip().upper() not in {"M²", "M2"}
                   for index in quantity_headers)
        )
        if not compact and not detailed:
            raise ValueError("warehouse workbook is missing the detailed inventory headers")
        mapping = _catalog_map(catalog)
        totals: dict[tuple[str, str], Decimal] = defaultdict(lambda: Decimal("0.00"))
        display_refs: dict[tuple[str, str], str] = {}
        source_refs: set[tuple[str, str]] = set()
        diagnostics = []
        source_rows = 0
        for row_number, source in enumerate(
                _bounded_rows(sheet, min_row=3, values_only=True), start=3):
            values = list(source)
            if len(values) < (5 if compact else 15):
                values.extend([None] * ((5 if compact else 15) - len(values)))
            brand = str(values[0] or "").strip().upper()
            category = str(values[1] or "").strip().upper()
            if not brand.startswith("TARRAGONA") or "CERAMICO" not in category:
                continue
            source_rows += 1
            raw_ref = str(values[2] or "").strip()
            normalized_ref = _norm(raw_ref)
            canonical = mapping.get(normalized_ref)
            if not normalized_ref:
                diagnostics.append({"source_row_ref": f"row-{row_number}",
                                    "severity": "error", "message": "missing warehouse reference"})
                continue
            identity = ("catalog", canonical) if canonical else ("source", normalized_ref)
            display_refs.setdefault(identity, canonical or raw_ref)
            try:
                final = _decimal(values[4] if compact else values[14], "final M2")
                operands = () if compact else (values[5], values[8], values[11])
                if operands and all(value not in (None, "") for value in operands):
                    initial, incoming, outgoing = (
                        _decimal(values[5], "initial M2"),
                        _decimal(values[8], "incoming M2"),
                        _decimal(values[11], "outgoing M2"))
                    if abs(initial + incoming - outgoing - final) > Q2:
                        diagnostics.append({"source_row_ref": f"row-{row_number}",
                                            "severity": "error",
                                            "message": "warehouse arithmetic mismatch"})
                        continue
                if final < 0:
                    raise ValueError("negative final M2")
            except ValueError as exc:
                diagnostics.append({"source_row_ref": f"row-{row_number}",
                                    "severity": "error", "message": str(exc)})
                continue
            totals[identity] += final
            source_refs.add((normalized_ref, raw_ref))
        rows = [{"product_ref": display_refs[key], "m2": _q2(value)}
                for key, value in sorted(totals.items())]
        return {
            "rows": rows,
            "source_refs": [
                {"source_family": "warehouse",
                 "normalized_identity": normalized,
                 "raw_ref": raw}
                for normalized, raw in sorted(source_refs)
            ],
            "diagnostics": diagnostics,
            "source_rows": source_rows,
        }
    finally:
        workbook.close()


def parse_sales_xlsx(file_bytes: bytes, *, catalog: Sequence[dict]) -> dict:
    """Derive one 90-day velocity and maximum ISO-week total per product."""
    workbook = _open_xlsx(file_bytes)
    try:
        formula_workbook = _open_xlsx(file_bytes, data_only=False)
        try:
            return _parse_sales_workbooks(workbook, formula_workbook, catalog=catalog)
        finally:
            formula_workbook.close()
    finally:
        workbook.close()


def _parse_sales_workbooks(workbook, formula_workbook, *, catalog: Sequence[dict]) -> dict:
    """Parse cached values while retaining source formula visibility."""
    sheet = workbook["VENTAS"] if "VENTAS" in workbook.sheetnames else workbook.active
    formula_sheet = (formula_workbook["VENTAS"] if "VENTAS" in formula_workbook.sheetnames
                     else formula_workbook.active)
    headers = [str(value or "").strip().upper()
               for value in next(_bounded_rows(sheet, min_row=1, max_row=1, values_only=True))]
    required = {"FECHA", "REFERENCIA", "MT2"}
    if not required.issubset(headers):
        raise ValueError("sales workbook is missing required headers")
    indexes = {name: headers.index(name) for name in required}
    mapping = _catalog_map(catalog)
    records = []
    diagnostics = []
    sources = _bounded_rows(sheet, min_row=2, values_only=True)
    formula_sources = _bounded_rows(formula_sheet, min_row=2, values_only=False)
    for row_number, (source, formula_source) in enumerate(
            zip(sources, formula_sources), start=2):
        values = list(source)
        if not any(value not in (None, "") for value in values):
            continue
        mt2_index = indexes["MT2"]
        mt2_value = values[mt2_index] if mt2_index < len(values) else None
        formula_cell = formula_source[mt2_index] if mt2_index < len(formula_source) else None
        if formula_cell is not None and formula_cell.data_type == "f":
            diagnostics.append({"source_row_ref": f"row-{row_number}",
                                "severity": "error",
                                "message": "formula MT2 cannot be evaluated"})
            continue
        if mt2_value is None or (isinstance(mt2_value, str) and not mt2_value.strip()):
            continue
        try:
            sold_on = _excel_date(values[indexes["FECHA"]])
            raw_identity = values[indexes["REFERENCIA"]]
            canonical = (mapping.get(_norm(raw_identity))
                         if raw_identity is not None else None)
            identity = canonical or _identity(raw_identity, mapping)
            sold = _decimal(mt2_value, "MT2")
            if sold_on is None or identity is None:
                raise ValueError("invalid sale date or reference")
        except (IndexError, ValueError) as exc:
            diagnostics.append({"source_row_ref": f"row-{row_number}",
                                "severity": "error", "message": str(exc)})
            continue
        if sold > 0:
            records.append((sold_on, identity, sold, canonical is not None))
    matched_records = [record for record in records if record[3]]
    if not matched_records:
        raise ValueError("sales workbook contains no positive catalog-matched sales")
    as_of = max(record[0] for record in matched_records)
    window_start = as_of - timedelta(days=89)
    by_product: dict[str, list[tuple[date, Decimal]]] = defaultdict(list)
    for sold_on, identity, sold, catalog_matched in records:
        if ((catalog_matched and window_start <= sold_on <= as_of)
                or (not catalog_matched and sold_on >= window_start)):
            by_product[identity].append((sold_on, sold))
    rows = []
    for identity in sorted(by_product):
        product_records = by_product[identity]
        weekly: dict[tuple[int, int], Decimal] = defaultdict(lambda: Decimal("0.00"))
        total = Decimal("0.00")
        for sold_on, sold in product_records:
            iso = sold_on.isocalendar()
            weekly[(iso.year, iso.week)] += sold
            total += sold
        rows.append({"product_ref": identity,
                     "daily_velocity": _q2(total / Decimal(90)),
                     "peak_weekly_m2": _q2(max(weekly.values()))})
    return {"as_of": as_of, "rows": rows, "diagnostics": diagnostics}


def _siesa_item(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value).strip())
    except ArithmeticError:
        return None
    return int(number) if number == number.to_integral_value() else None


def parse_siesa_xlsx(file_bytes: bytes, *, catalog: Sequence[dict]) -> dict:
    """Normalize and aggregate SIESA lots without replacing source availability."""
    workbook = _open_xlsx(file_bytes)
    try:
        sheet = workbook.active
        return _parse_siesa_values(_bounded_rows(sheet, values_only=True), catalog=catalog)
    finally:
        workbook.close()


def parse_siesa_xls(file_bytes: bytes, *, catalog: Sequence[dict]) -> dict:
    """Read the bounded legacy OLE/BIFF8 export used only by SIESA."""
    if (not file_bytes or len(file_bytes) > _MAX_XLS_BYTES
            or not file_bytes.startswith(b"\xd0\xcf\x11\xe0")):
        raise ValueError("invalid SIESA XLS workbook")
    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=file_bytes, on_demand=True)
        if len(workbook.sheet_names()) != 1 or len(workbook.sheet_names()) > _MAX_XLS_SHEETS:
            raise ValueError
        sheet = workbook.sheet_by_index(0)
        if (sheet.nrows > _MAX_WORKSHEET_ROWS or sheet.ncols > _MAX_WORKSHEET_COLUMNS
                or sheet.nrows * sheet.ncols > _MAX_WORKSHEET_CELLS):
            raise ValueError
        return _parse_siesa_values(
            (sheet.row_values(index) for index in range(sheet.nrows)), catalog=catalog)
    except (xlrd.XLRDError, OSError, ValueError, IndexError) as exc:
        raise ValueError("invalid SIESA XLS workbook") from exc
    finally:
        if "workbook" in locals():
            workbook.release_resources()


def _parse_siesa_values(source_rows, *, catalog: Sequence[dict]) -> dict:
    iterator = iter(source_rows)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration as exc:
        raise ValueError("SIESA workbook is missing required headers") from exc
    required = ["Item", "Existencia", "Cant. comprometida", "Cant. disponible"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("SIESA workbook is missing required headers: " + ", ".join(missing))
    indexes = {name: headers.index(name) for name in required}
    description_index = headers.index("Desc. item") if "Desc. item" in headers else None
    item_map: dict[int, str] = {}
    item_collisions: set[int] = set()
    name_map = _catalog_map(catalog)
    for product in catalog:
        item = _siesa_item(product.get("siesa_item"))
        if item is None:
            continue
        canonical = str(product["sku"])
        prior = item_map.get(item)
        if prior is not None and prior != canonical:
            item_collisions.add(item)
        else:
            item_map[item] = canonical
    for item in item_collisions:
        item_map.pop(item, None)
    totals: dict[str, dict[str, Decimal]] = defaultdict(
        lambda: {"available": Decimal("0.00"), "committed": Decimal("0.00")})
    source_refs: set[tuple[str, str]] = set()
    diagnostics = []
    for row_number, source in enumerate(iterator, start=2):
        values = list(source)
        if not any(value not in (None, "") for value in values):
            continue
        raw_item = values[indexes["Item"]] if indexes["Item"] < len(values) else None
        description = (values[description_index] if description_index is not None
                       and description_index < len(values) else None)
        quantity_values = [values[indexes[name]] if indexes[name] < len(values) else None
                           for name in required[1:]]
        if (str(description or "").strip().upper() == "TOTAL"
                or (not str(description or "").strip()
                    and all(value in (None, "") for value in quantity_values))):
            continue
        try:
            existence = _decimal(values[indexes["Existencia"]], "Existencia")
            committed = _decimal(values[indexes["Cant. comprometida"]], "Cant. comprometida")
            available = _decimal(values[indexes["Cant. disponible"]], "Cant. disponible")
        except (IndexError, ValueError) as exc:
            diagnostics.append({"source_row_ref": f"row-{row_number}",
                                "severity": "error", "message": str(exc)})
            continue
        if min(existence, committed, available) < 0:
            diagnostics.append({"source_row_ref": f"row-{row_number}", "severity": "error",
                                "message": "negative required SIESA quantity"})
            continue
        item = _siesa_item(raw_item)
        identity = item_map.get(item) if item is not None else None
        identity = identity or _identity(description, name_map)
        if identity is None:
            identity = str(raw_item or "").strip() or "unidentified SIESA item"
        if abs(existence - committed - available) > Q2:
            diagnostics.append({"source_row_ref": f"row-{row_number}", "severity": "warning",
                                "message": "Existencia - Cant. comprometida differs from Cant. disponible"})
        totals[identity]["available"] += available
        totals[identity]["committed"] += committed
        raw_ref = str(description or raw_item or "").strip()
        if raw_ref:
            source_refs.add((_norm(identity), raw_ref))
    rows = [{"product_ref": key, "available_m2": _q2(value["available"]),
             "committed_m2": _q2(value["committed"])}
            for key, value in sorted(totals.items())]
    return {
        "rows": rows,
        "source_refs": [
            {"source_family": "siesa_availability",
             "normalized_identity": normalized,
             "raw_ref": raw}
            for normalized, raw in sorted(source_refs)
        ],
        "diagnostics": diagnostics,
    }


def _month_date(value: str) -> date:
    return datetime.strptime(value.replace(",2026", ", 2026"), "%b %d, %Y").date()


def parse_production_schedule_text(text: str, *, catalog: Sequence[dict],
                                   evidence_as_of: date) -> dict:
    """Parse the text layer of Ashley's PLAN_DE_PRODUCCION report.

    TERMINADO is trusted completion evidence, but never SIESA availability.
    Delivery dates remain estimates; no start/actual-ready date is invented.
    """
    mapping = _catalog_map(catalog)
    rows: list[dict] = []
    unmatched: list[str] = []
    statuses = {
        "TERMINADO": ("completed", True),
        "PRODUCCIÓN": ("in_progress", False),
        "PENDIENTE": ("scheduled", False),
    }
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if "52X52" not in line:
            continue
        prefix, suffix = line.split("52X52", 1)
        identity = re.match(
            r"^(P[12]-\d{5})\s+\d+\s+\d+\s+[\d.]+\s+\d+\s+(.+?)\s*$",
            prefix)
        quantities = re.search(
            r"([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+"
            r"(TERMINADO|PRODUCCIÓN|PENDIENTE)$", suffix)
        if identity is None or quantities is None:
            continue
        production_ref, raw_product = identity.groups()
        _planned_total, _actual_total, programmed_export, real_export, raw_status = quantities.groups()
        delivery_dates = re.findall(r"[A-Z][a-z]{2} \d{1,2},\s?2026", suffix)
        if not delivery_dates:
            continue
        # The second full date is the updated delivery estimate when present;
        # malformed/CAMBIO rows retain the initial estimate rather than inventing one.
        active_delivery = delivery_dates[1] if len(delivery_dates) > 1 else delivery_dates[0]
        normalized = _norm(raw_product)
        canonical = mapping.get(normalized)
        if canonical is None:
            if not normalized:
                continue
            if normalized not in unmatched:
                unmatched.append(normalized)
        status, completion_confirmed = statuses[raw_status]
        export_m2 = real_export if completion_confirmed else programmed_export
        rows.append({
            "product_ref": canonical or normalized,
            "m2": format(_q2(export_m2.replace(",", "")), "f"),
            "scheduled_start": None,
            "status": status,
            "production_ref": production_ref,
            "estimated_ready_date": _month_date(active_delivery),
            "actual_ready_date": None,
            "completion_confirmed": completion_confirmed,
            "can_add_more": False,
            "evidence_as_of": evidence_as_of,
        })
    return {"rows": rows, "unmatched_products": unmatched,
            "authority": "production_context_not_siesa_availability"}


def parse_production_schedule_pdf(pdf_bytes: bytes, *, catalog: Sequence[dict],
                                  evidence_as_of: date) -> dict:
    """Extract a text-layer PDF locally, then apply the deterministic parser."""
    import pdfplumber
    from pdfminer.pdfexceptions import PDFException
    from pdfplumber.utils.exceptions import PdfminerException

    try:
        with pdfplumber.open(BytesIO(pdf_bytes)) as document:
            if len(document.pages) > _MAX_PDF_PAGES:
                raise ValueError
            started = _monotonic()
            parts = []
            tables = []
            text_bytes = 0
            for page in document.pages:
                if _monotonic() - started > _MAX_PDF_SECONDS:
                    raise ValueError
                part = page.extract_text() or ""
                if _monotonic() - started > _MAX_PDF_SECONDS:
                    raise ValueError
                text_bytes += len(part.encode("utf-8"))
                if text_bytes > _MAX_PDF_TEXT_BYTES:
                    raise ValueError
                parts.append(part)
                extract_tables = getattr(page, "extract_tables", None)
                if callable(extract_tables):
                    tables.extend(extract_tables() or [])
            text = "\n".join(parts)
            if not text.strip():
                raise ValueError
    except (PdfminerException, PDFException, OSError, ValueError) as exc:
        raise ValueError("invalid production PDF") from exc
    structured = _parse_production_schedule_tables(
        tables, catalog=catalog, evidence_as_of=evidence_as_of)
    if structured is not None:
        return structured
    return parse_production_schedule_text(text, catalog=catalog, evidence_as_of=evidence_as_of)


_SPANISH_MONTHS = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
                   "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
                   "septiembre": 9, "octubre": 10, "noviembre": 11,
                   "diciembre": 12}


def _spanish_date(value) -> date | None:
    match = re.search(r"(\d{1,2}) de ([a-záéíóú]+) de (\d{4})",
                      str(value or "").lower())
    if not match:
        return None
    day, month, year = match.groups()
    normalized_month = unicodedata.normalize("NFKD", month).encode("ascii", "ignore").decode()
    month_number = _SPANISH_MONTHS.get(normalized_month)
    return date(int(year), month_number, int(day)) if month_number else None


def _report_m2(value) -> Decimal:
    text = str(value or "").strip()
    if not text:
        return Decimal("0.00")
    return _q2(text.replace(".", "").replace(",", "."))


def _parse_production_schedule_tables(tables, *, catalog: Sequence[dict], evidence_as_of: date):
    mapping = _catalog_map(catalog)
    candidates = []
    for table in tables:
        for source in table:
            values = list(source)
            if (len(values) < 18 or str(values[2] or "").strip() != "52X52"
                    or re.fullmatch(r"P[12]-\d{5}", str(values[6] or "").strip()) is None):
                continue
            candidates.append(values)
    if not candidates:
        return None
    rows = []
    unmatched = []
    unmatched_emitted = set()
    for values in candidates:
        programmed = _report_m2(values[15])
        real = _report_m2(values[17])
        amount = real if real > 0 else programmed
        if amount <= 0:
            continue
        normalized = _norm(values[9])
        canonical = mapping.get(normalized)
        if canonical is None:
            if not normalized or normalized in unmatched_emitted:
                continue
            unmatched_emitted.add(normalized)
            unmatched.append(normalized)
        ready = _spanish_date(values[0])
        rows.append({
            "product_ref": canonical or normalized,
            "m2": format(amount, "f"),
            "scheduled_start": None,
            "status": "in_progress" if ready and ready <= evidence_as_of else "scheduled",
            "production_ref": str(values[6]).strip(),
            "estimated_ready_date": ready,
            "actual_ready_date": None,
            "completion_confirmed": False,
            "can_add_more": False,
            "evidence_as_of": evidence_as_of,
        })
    return {"rows": rows, "unmatched_products": unmatched,
            "authority": "production_context_not_siesa_availability",
            "source_rows_inspected": len(candidates),
            "omitted_rows": len(candidates) - len(rows)}


def _excel_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        return date(1899, 12, 30) + timedelta(days=int(value))
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def parse_scheduled_dispatch_xlsx(file_bytes: bytes, *, catalog: Sequence[dict]) -> dict:
    """Parse dispatch bookings as scheduled context, never transit evidence."""
    workbook = _open_xlsx(file_bytes)
    try:
        formula_workbook = _open_xlsx(file_bytes, data_only=False)
        try:
            return _parse_dispatch_workbooks(workbook, formula_workbook, catalog=catalog)
        finally:
            formula_workbook.close()
    finally:
        workbook.close()


def _parse_dispatch_workbooks(workbook, formula_workbook, *, catalog: Sequence[dict]) -> dict:
    """Parse cached values while rejecting formulas in the M2 source column."""
    sheet_name = "DETALLE DE DESPACHO - GUATEMALA"
    if sheet_name not in workbook.sheetnames or sheet_name not in formula_workbook.sheetnames:
        raise ValueError("dispatch workbook is missing required sheet")
    sheet = workbook[sheet_name]
    formula_sheet = formula_workbook[sheet_name]
    mapping = _catalog_map(catalog)
    current: dict = {}
    orders: dict[str, dict] = {}
    unmatched: list[str] = []
    source_lines = 0
    sources = _bounded_rows(sheet, values_only=True)
    formula_sources = _bounded_rows(formula_sheet, values_only=False)
    for row_number, (values, formula_values) in enumerate(
            zip(sources, formula_sources), start=1):
        if row_number < 4:
            continue
        values = list(values)
        while len(values) < 9:
            values.append(None)
        raw_product = values[7]
        amount = values[8]
        if not raw_product or str(raw_product).strip().upper() == "TOTAL":
            continue
        formula_m2 = formula_values[8] if len(formula_values) > 8 else None
        if formula_m2 is not None and formula_m2.data_type == "f":
            raise ValueError("invalid dispatch workbook")
        if amount in (None, ""):
            continue
        source_lines += 1
        for index, field in ((0, "purchase_order"), (1, "pedido"),
                             (2, "containers"), (3, "etd"),
                             (4, "eta"), (5, "booking")):
            if values[index] not in (None, ""):
                current[field] = values[index]
        m2 = _decimal(amount, "dispatch M2")
        if m2 <= 0:
            continue
        canonical = mapping.get(_norm(raw_product))
        if canonical is None:
            normalized = _norm(raw_product)
            if normalized and normalized not in unmatched:
                unmatched.append(normalized)
            continue
        order_ref = str(current.get("purchase_order") or current.get("pedido") or "").strip()
        if not order_ref:
            continue
        if order_ref not in orders:
            etd = _excel_date(current.get("etd"))
            eta = _excel_date(current.get("eta"))
            quality_flags = []
            if etd is not None and eta is not None and eta <= etd:
                quality_flags.append("eta_not_after_etd")
            orders[order_ref] = {
                "purchase_order": order_ref,
                "pedido": str(current.get("pedido") or ""),
                "containers": int(current.get("containers") or 0),
                "booking": str(current.get("booking") or ""),
                "etd_tentative": etd,
                "eta_tentative": eta,
                "departure_state": "scheduled_not_departed",
                "quality_flags": quality_flags,
                "lines": [],
            }
        orders[order_ref]["lines"].append({
            "product_ref": canonical, "m2": _q2(m2),
            "raw_product": str(raw_product).strip()})
    return {"authority": "scheduled_tentative", "orders": list(orders.values()),
            "unmatched_products": unmatched, "source_lines": source_lines}
